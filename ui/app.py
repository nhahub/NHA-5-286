"""Gradio 6 dashboard for AI assisted test automation (multi-page, themed; backend wiring unchanged)."""

from __future__ import annotations
import html
import os
import re
import textwrap
import threading
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from xml.sax.saxutils import escape
import gradio as gr

from uuid import uuid4
from langgraph.types import Command
from graph.build_graph import build_graph
from ui import design as D
from ui import history as H
from ui import views as V
from ui.design import CSS, THEMES, theme_css, request_layout_css
from llm.settings import (
    LLMSettings, check_static, effective_key, register_run_credentials, release_run_credentials, validate_settings,
)
from llm.run_auth import AuthConfig, build_auth, public_request_fields, register_run_auth, release_run_auth
from preflight import run_preflight



ROOT = Path(__file__).resolve().parents[1]

WORKFLOW = build_graph()


EXPORT_DIR = ROOT / "runs" / "mock_reports"



STAGES = ["Intake", "Planner", "Code Generator", "Validation", "Execution", "Evidence Collection", "Analyzer", "Fix Proposal", "Human Approval", "Apply Fix (Healing)", "Report Generation"]

STATUS_LABELS = {"waiting for approval": "Waiting for Approval"}

STAGE_STATUSES = {"not run", "running", "waiting for approval", "completed", "failed", "skipped"}





PAGES = [label for label, _ in D.NAV_ORDER]

NAV_CHOICES = [(label, label) for label in PAGES]


def toggle_sidebar(collapsed: bool, page: str):
    """Switch between the full menu and a narrow icon-only rail; content reflows via --sb-w."""
    collapsed = not collapsed
    return (collapsed, D.collapsed_css() if collapsed else "", gr.update(value=page),
            gr.update(value="▶" if collapsed else "◀ Collapse"))


def progress_visibility(show_progress: bool) -> tuple[gr.update, str]:

    shown = bool(show_progress)

    return gr.update(visible=shown), request_layout_css(shown)





def render_progress(statuses: dict[str, str] | None = None, notes: dict[str, str] | None = None) -> str:

    statuses = statuses or {}

    notes = notes or {}

    markup = []

    for stage in STAGES:

        raw_status = str(statuses.get(stage, "Not run")).strip().lower()

        status = raw_status if raw_status in STAGE_STATUSES else "not run"

        markup.append(

            f"<div class='step status-{status.replace(' ', '-')}'><span class='step-dot'></span>"

            f"<span class='step-name'>{html.escape(stage)}</span><span class='step-status'>{STATUS_LABELS.get(status, status.title())}{html.escape(' · ' + notes[stage]) if notes.get(stage) else ''}</span></div>"

        )

    return "".join(markup)





METRIC_STYLE = {"Total Scenarios": ("", "results"), "Passed": ("ok", "check"), "Failed": ("bad", "x"), "Skipped": ("", "alert"),
                "Execution Time": ("info", "clock"), "Execution Time (simulated)": ("info", "clock"), "Healing Attempts": ("warn", "zap")}


def metric_card(value: str, label: str) -> str:
    tone, ico = METRIC_STYLE.get(label, ("", "activity"))
    return D.kpi(label, value, "", tone if value not in ("—", "0") or tone in ("", "info") else "", ico)


def set_theme(theme: str) -> tuple[str, str]:

    return theme_css(theme), theme





def navigate(page: str) -> list:
    return [gr.update(visible=page == name) for name in PAGES]


def topbar_html(page: str, use_mock: bool) -> str:
    mocking = use_mock is True or str(use_mock).strip().lower() in {"true", "1", "yes", "on"}
    mode = D.badge("Mock mode · simulated", "violet") if mocking else D.badge("Live mode · real agents", "info")
    return f"<div class='crumb'>Bugle <span>/</span> <b>{html.escape(page)}</b> {mode}</div>"


def session_mode_label(use_mock: bool) -> str:
    mocking_enabled = use_mock is True or str(use_mock).strip().lower() in {"true", "1", "yes", "on"}
    mode = "Mock mode (simulated, no real LLM)" if mocking_enabled else "Live mode (real agents only)"
    return f"<div class='session-mode'>Session mode<b>{mode}</b></div>"


def save_llm_settings(provider: str, model: str, api_key: str) -> tuple[str, dict]:

    """Validate the credentials with a real provider call, then keep them for this UI session only."""

    ok, message = validate_settings(provider, model or "", api_key or "")

    # Session-scoped only. Never expose this dict in report content, logs, or component values.
    state = {"provider": provider, "model": (model or "").strip(), "api_key": api_key or "", "validated": ok}

    if not ok:
        return f"**Not validated — live mode is unavailable with these settings.**\n\n{message}", state

    return f"**Validated.** {message} The API key stays masked, session-only, and is not exported.", state





def scenario_details(run: dict | None, evt: gr.SelectData) -> str:

    if not run or evt.index is None:

        return "Select a scenario row to inspect its details."

    index = evt.index[0] if isinstance(evt.index, (tuple, list)) else evt.index

    if not isinstance(index, int) or index >= len(run["scenarios"]):

        return "Select a scenario row to inspect its details."

    s = run["scenarios"][index]

    fields = [("Test Steps", "\n".join(f"{i}. {step}" for i, step in enumerate(s["Test Steps"], 1))),

              ("Expected Result", s["Expected Result"]), ("Actual Result", s["Actual Result"]),

              ("Error Message", s["Error Message"]), ("Screenshot / Evidence", s["Evidence Path"]),

              ("Analysis", s["Analysis"]), ("Proposed Fix", s["Proposed Fix"])]

    return "\n\n".join(f"### {name}\n{value}" for name, value in fields)





def _report_notice(run: dict) -> str:
    if run["simulated"]:
        return "SIMULATED mock run: agents, browser execution, and fixes are mocks. This is not a real test result."
    return "Live run."


def _report_facts(run: dict) -> list[tuple[str, str]]:
    models = ", ".join(f"{k}: {v}" for k, v in run["models_used"].items()) or "none recorded"
    return [("Run ID", run["run_id"]), ("Mode", run["mode"]), ("Created", run["created_at"]),
            ("Application URL", run["url"]), ("Requirements", run["requirements"]),
            ("Browser", f"{run['browser']} / {run['execution_mode']}"), ("Terminal status", run["terminal_status"]),
            ("Total scenarios", str(len(run["scenarios"]))), ("Passed", str(run["passed"])), ("Failed", str(run["failed"])),
            ("Skipped / not run", str(run["skipped"])), ("Execution time", run["execution_time"] + (" (simulated)" if run["simulated"] and run["execution_time"][0].isdigit() else "")),
            ("Healing attempts", f"{run['healing_attempts']} of {MAX_HEALING}"), ("Models used", models),
            ("Executed nodes", " > ".join(run["trail"]) or "none")]


def preview_report(run: dict | None) -> str:
    if not run:
        return "No run is selected. Start a test first."
    lines = [f"# Test Run Report — {run['run_id']}", f"> {_report_notice(run)}", ""]
    lines += [f"**{label}:** {value}  " for label, value in _report_facts(run)]
    lines += ["", "## Scenario results", "| ID | Scenario | Status | Duration | Error summary |", "|---|---|---|---|---|"]
    lines.extend(f"| {s['Scenario ID']} | {s['Scenario Name']} | {s['Status']} | {s['Duration']} | {s['Error Summary']} |"
                 for s in run["scenarios"])
    for s in run["scenarios"]:
        lines += ["", f"### {s['Scenario ID']} — {s['Scenario Name']} ({s['Status']})",
                  f"**Evidence**  \n{s['Evidence Path']}", f"**Analysis**  \n{s['Analysis']}", f"**Proposed fix**  \n{s['Proposed Fix']}"]
    if run["workflow_error"]:
        lines += ["", f"**Workflow error:** {run['workflow_error']}"]
    return "\n".join(lines)


def _safe(value: object) -> str:
    return str(value if value is not None else "")


def _plain_lines(run: dict) -> list[str]:
    """Report text without Markdown markup, shared by the PDF and Word writers."""
    lines = [f"Test Run Report - {run['run_id']}", _report_notice(run), ""]
    lines += [f"{label}: {value}" for label, value in _report_facts(run)]
    lines += ["", "Scenario results"]
    for s in run["scenarios"]:
        lines += [f"{s['Scenario ID']} | {s['Scenario Name']} | {s['Status']} | {s['Duration']} | {s['Error Summary']}",
                  "Evidence: " + s["Evidence Path"].replace("\n", "; "),
                  "Analysis: " + re.sub(r"\s+", " ", s["Analysis"]),
                  "Proposed fix: " + re.sub(r"[`\n]+", " ", s["Proposed Fix"])]
    if run["workflow_error"]:
        lines += ["", "Workflow error: " + re.sub(r"[*`]", "", run["workflow_error"])]
    return lines


def _write_pdf(path: Path, run: dict) -> None:
    """Write a small standards-compliant, text-only PDF without a runtime dependency."""
    text_lines: list[str] = []
    for line in _plain_lines(run):
        text_lines += textwrap.wrap(line, 90) or [""]
    pages = [text_lines[i:i + 48] for i in range(0, len(text_lines), 48)] or [[]]
    objects: list[bytes] = []
    page_ids = [4 + 2 * i for i in range(len(pages))]
    font_id = 3
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(("<< /Type /Pages /Kids [" + " ".join(f"{i} 0 R" for i in page_ids) + f"] /Count {len(pages)} >>").encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for i, page in enumerate(pages):
        page_id, content_id = 4 + 2*i, 5 + 2*i
        stream = ["BT", "/F1 10 Tf", "50 790 Td", "14 TL"]
        for line in page:
            ascii_line = line.encode("ascii", "replace").decode().replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            stream += [f"({ascii_line}) Tj", "T*"]
        stream.append("ET")
        payload = "\n".join(stream).encode("ascii")
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] /Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_id} 0 R >>".encode())
        objects.append(b"<< /Length " + str(len(payload)).encode() + b" >>\nstream\n" + payload + b"\nendstream")
    data = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(data)); data += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(data); data += f"xref\n0 {len(objects)+1}\n0000000000 65535 f \n".encode()
    for offset in offsets[1:]: data += f"{offset:010d} 00000 n \n".encode()
    data += f"trailer\n<< /Size {len(objects)+1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    path.write_bytes(data)


def _write_docx(path: Path, run: dict) -> None:
    paragraphs = _plain_lines(run)
    body = "".join(f'<w:p><w:r><w:t xml:space="preserve">{escape(p)}</w:t></w:r></w:p>' for p in paragraphs)
    document = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>{body}<w:sectPr><w:pgSz w:w="12240" w:h="15840"/><w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440"/></w:sectPr></w:body></w:document>'''
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as doc:
        doc.writestr("[Content_Types].xml", '<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
        doc.writestr("_rels/.rels", '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
        doc.writestr("word/document.xml", document)


def export_report(run: dict | None, format_name: str) -> tuple[str, str | None]:
    if not run:
        return "Run a test before exporting.", None
    if format_name not in {"PDF", "Word (.docx)", "Excel (.xlsx)"}:
        return "Choose PDF, Word, or Excel before exporting.", None
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    stem = re.sub(r"[^A-Za-z0-9_-]", "_", run["run_id"])
    if format_name == "Excel (.xlsx)":
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            return "Excel export requires openpyxl. Install it with `pip install openpyxl`.", None
        path = EXPORT_DIR / f"{stem}.xlsx"
        book = Workbook(); summary = book.active; summary.title = "Summary"
        summary.append(["Field", "Value"])
        summary.append(["Notice", _report_notice(run)])
        for row in _report_facts(run): summary.append(list(row))
        scenarios = book.create_sheet("Scenarios"); scenarios.append(["Scenario ID", "Scenario Name", "Status", "Duration", "Error Summary"])
        for s in run["scenarios"]: scenarios.append([s[k] for k in ["Scenario ID", "Scenario Name", "Status", "Duration", "Error Summary"]])
        failures = book.create_sheet("Failure Analysis"); failures.append(["Scenario ID", "Scenario Name", "Analysis", "Error Message", "Evidence", "Proposed Fix"])
        for s in run["scenarios"]: failures.append([s["Scenario ID"], s["Scenario Name"], s["Analysis"], s["Error Message"], s["Evidence Path"], s["Proposed Fix"]])
        for sheet in book.worksheets:
            sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="315EA8")
            for col in sheet.columns:
                width = min(60, max(14, max(len(_safe(c.value)) for c in col) + 2)); sheet.column_dimensions[col[0].column_letter].width = width
        book.save(path)
    elif format_name == "Word (.docx)":
        path = EXPORT_DIR / f"{stem}.docx"; _write_docx(path, run)
    else:
        path = EXPORT_DIR / f"{stem}.pdf"; _write_pdf(path, run)
    return f"Exported {format_name} report for {run['run_id']}.", str(path)


# ----------------------------------------------------------------------------- workflow
MAX_HEALING = 2

NODE_STAGE = {
    "intake": "Intake", "planner": "Planner", "codegen": "Code Generator", "validate": "Validation",
    "execute": "Execution", "analyze": "Analyzer", "fix_proposal": "Fix Proposal",
    "fix_approval": "Human Approval", "apply_fix": "Apply Fix (Healing)", "report": "Report Generation",
}


class StageTracker:
    """Per-run stage status derived only from LangGraph node events (no static previews)."""

    def __init__(self) -> None:
        self.status = {stage: "Not run" for stage in STAGES}
        self.count = {stage: 0 for stage in STAGES}
        self.note: dict[str, str] = {}
        self.errors_seen = 0
        self.history: list[dict] = []  # every pass through every stage, kept across retries/reruns

    ACTIVE = ("Running", "Waiting for Approval")

    def _set(self, stage: str, status: str, note: str | None = None) -> None:
        self.status[stage] = status
        if note is not None:
            self.note[stage] = note
        if status == "Skipped" or not self.count[stage]:
            return
        entry = next((h for h in reversed(self.history) if h["stage"] == stage and h["pass"] == self.count[stage]), None)
        if entry is None:
            entry = {"stage": stage, "pass": self.count[stage], "status": status, "note": ""}
            self.history.append(entry)
        entry["status"] = status
        if note is not None:
            entry["note"] = note

    def begin(self, node: str, note: str | None = None) -> None:
        stage = NODE_STAGE.get(node)
        if stage is None:
            return
        if self.status[stage] not in self.ACTIVE:
            self.count[stage] += 1  # a revisit (re-validation, rerun, second analysis) is a new pass
        waiting = note == "awaiting decision"
        self._set(stage, "Waiting for Approval" if waiting else "Running", note or "")
        if node == "execute":
            if self.status["Evidence Collection"] not in self.ACTIVE:
                self.count["Evidence Collection"] += 1
            self._set("Evidence Collection", "Running", "")

    def finish(self, node: str, update: dict) -> None:
        stage = NODE_STAGE.get(node)
        if stage is None:
            return
        if self.status[stage] not in self.ACTIVE:
            self.begin(node)
        terminal = update.get("terminal_status")
        failed, note = False, ""
        if node == "intake" and terminal in ("BLOCKED", "ERROR"):
            failed, note = True, terminal
        elif node == "validate" and update.get("validation_status") == "invalid":
            failed, note = True, "invalid"
        elif node == "execute":
            failed, note = terminal != "PASS", terminal or ""
        elif node == "apply_fix":
            failed = terminal == "FAIL"
            note = "not applied" if failed else f"attempt {update.get('healing_retry_count', '?')} of {MAX_HEALING}"
        elif node == "fix_approval":
            note = str(update.get("approval_decision") or "")
        elif node == "analyze":
            note = str(update.get("failure_type") or "")
        errors = update.get("errors")
        if isinstance(errors, list):
            if len(errors) > self.errors_seen:
                failed, note = True, (note or "error")
            self.errors_seen = max(self.errors_seen, len(errors))
        self._set(stage, "Failed" if failed else "Completed", note)
        if node == "execute":
            has_evidence = bool(update.get("execution_evidence"))
            self._set("Evidence Collection", "Completed" if has_evidence else "Skipped",
                      "" if has_evidence else "no evidence")

    def fail_running(self, reason: str) -> None:
        for stage, status in self.status.items():
            if status in self.ACTIVE:
                self._set(stage, "Failed", reason)

    def finalize(self) -> None:
        for stage, status in self.status.items():
            if status == "Not run":
                self.status[stage] = "Skipped"

    def view(self) -> tuple[dict[str, str], dict[str, str]]:
        notes = {}
        for stage in STAGES:
            parts = ["" if self.note.get(stage) == "awaiting decision" else self.note.get(stage, "")]
            if self.count[stage] > 1:
                parts.append(f"×{self.count[stage]}")
            notes[stage] = " ".join(p for p in parts if p)
        return dict(self.status), notes


RUNS: dict[str, dict] = {}  # thread_id -> UI-side run info; graph state itself lives in the checkpointer
_RUNS_LOCK = threading.Lock()
NOOP = gr.update()


def _config(thread_id: str) -> dict:
    return {"configurable": {"thread_id": thread_id}}


def release_run_secrets(thread_id: str) -> None:
    """Drop everything secret held for a run: the LLM key and the application login."""
    release_run_credentials(thread_id)
    release_run_auth(thread_id)


def _scrub(text: object, info: dict | None) -> str:
    """Remove any API key or application credential from a message before it reaches the UI."""
    from llm.errors import redact

    text = redact(str(text))
    key = (info or {}).get("secret") or ""
    return text.replace(key, "***") if key else text


def _pending_interrupt(thread_id: str) -> dict | None:
    """The approval request this run is actually paused on, read from the checkpoint."""
    try:
        snapshot = WORKFLOW.get_state(_config(thread_id))
    except Exception:
        return None
    for task in snapshot.tasks:
        for item in task.interrupts:
            return dict(item.value) if isinstance(item.value, dict) else {"message": str(item.value)}
    return None


def _trail_label(node: str, update: dict) -> str:
    detail = None
    if node == "execute":
        detail = update.get("terminal_status")
    elif node == "validate":
        detail = update.get("validation_status")
    elif node == "analyze":
        detail = update.get("failure_type")
    elif node == "fix_approval":
        detail = update.get("approval_decision")
    elif node == "apply_fix" and "healing_retry_count" in update:
        detail = f"attempt {update['healing_retry_count']}"
    return f"{node} ({detail})" if detail else node


def _advance(thread_id: str, graph_input):
    """Drive the real graph one event at a time, updating the tracker; yields after each event."""
    info = RUNS[thread_id]
    tracker = info["tracker"]
    try:
        # "tasks" events say which node has really started; "updates" events carry what it returned.
        for mode, data in WORKFLOW.stream(graph_input, config=_config(thread_id), stream_mode=["tasks", "updates"]):
            if mode == "tasks":
                if "input" in data:
                    tracker.begin(data["name"])
                elif data.get("interrupts"):
                    tracker.begin(data["name"], "awaiting decision")
            else:
                for node, update in data.items():
                    if node == "__interrupt__":
                        continue
                    update = update or {}
                    tracker.finish(node, update)
                    info["trail"].append(_trail_label(node, update))
            yield
    except GeneratorExit:  # the client disconnected or the event was cancelled mid-run
        info["error"] = "Run cancelled."
        tracker.fail_running("cancelled")
        info["ended"] = datetime.now(timezone.utc).isoformat(timespec="seconds")
        release_run_secrets(thread_id)
        info["secret"] = ""
        raise
    except NotImplementedError as exc:
        stage = next((s for s, st in tracker.status.items() if st in tracker.ACTIVE), "a stage")
        info["error"] = (f"The live workflow stopped at **{stage}**: {_scrub(exc, info)} "
                         "No mock data was substituted.")
        tracker.fail_running("not implemented")
    except Exception as exc:
        info["error"] = f"{type(exc).__name__}: {_scrub(exc, info)}"
        tracker.fail_running("error")
    info["paused"] = info.get("error") is None and _pending_interrupt(thread_id) is not None
    if info["paused"]:
        tracker.begin("fix_approval", "awaiting decision")
    if not info["paused"]:
        tracker.finalize()
        info["ended"] = datetime.now(timezone.utc).isoformat(timespec="seconds")
        release_run_secrets(thread_id)
        info["secret"] = ""
    _persist(thread_id)


def _persist(thread_id: str) -> None:
    """Save a summary of this run for the dashboards (UI-side bookkeeping only; never read back by the workflow)."""
    try:
        info, values = RUNS[thread_id], _snapshot_values(thread_id)
        run = build_run(thread_id)
        tests = (values.get("execution_result") or {}).get("tests") or []
        duration = sum(int(t.get("duration_ms") or 0) for t in tests) if tests else None
        status = "PENDING" if info.get("paused") else ("ERROR" if info.get("error") else run["terminal_status"])
        H.save_record({
            "id": thread_id, "run_id": run["run_id"], "source": "ui", "started": info["started"], "ended": info.get("ended"),
            "mock": info["mock"], "status": status, "url": info["meta"]["url"], "total": len(run["scenarios"]),
            "passed": run["passed"], "failed": run["failed"], "skipped": run["skipped"], "duration_ms": duration,
            "healing": run["healing_attempts"], "failure_type": (values.get("failure_analysis") or {}).get("failure_type"),
            "fix_proposed": bool(values.get("fix_proposal")), "decisions": list(info.get("decisions") or []),
            "trail": list(info["trail"]),
            "validation_failed": any(h["stage"] == "Validation" and h["status"] == "Failed" for h in info["tracker"].history),
        })
    except Exception:
        pass  # history is best-effort and must never interfere with the workflow


def _snapshot_values(thread_id: str) -> dict:
    try:
        return dict(WORKFLOW.get_state(_config(thread_id)).values or {})
    except Exception:
        return {}


def _status_text(thread_id: str) -> str:
    info, values = RUNS[thread_id], _snapshot_values(thread_id)
    healing = values.get("healing_retry_count", 0)
    if info.get("error"):
        return f"### Workflow stopped\n{info['error']}"
    if info.get("paused"):
        return f"Workflow paused for human approval. Healing attempts used so far: {healing} of {MAX_HEALING}."
    label = "SIMULATED" if info["mock"] else "LIVE"
    return (f"Workflow completed: {values.get('terminal_status', 'UNKNOWN')} ({label}). "
            f"Healing attempts used: {healing} of {MAX_HEALING}.")


def _progress_html(thread_id: str) -> str:
    tracker = RUNS[thread_id]["tracker"]
    markup = render_progress(*tracker.view())
    passes = [h for h in tracker.history if tracker.count[h["stage"]] > 1]
    if passes:  # only stages that were repeated need a history; the rows above show the latest pass
        items = "".join(f"<div>{html.escape(h['stage'])} #{h['pass']}: {html.escape(h['status'])}"
                        f"{html.escape(' · ' + h['note']) if h['note'] else ''}</div>" for h in passes)
        markup += f"<div class='muted' style='margin-top:10px'><b>Pass history</b>{items}</div>"
    return markup


def approval_ui(session: dict | None):
    """Approve/Reject panel for the run's real pending interrupt (or hidden when none)."""
    payload = _pending_interrupt(session["thread_id"]) if session and session.get("thread_id") in RUNS else None
    if not payload:
        return gr.update(value="", visible=False), gr.update(visible=False)
    values = _snapshot_values(session["thread_id"])
    analysis = values.get("failure_analysis") or {}
    used = values.get("healing_retry_count", 0)
    subtype = f" / {analysis['subtype']}" if analysis.get("subtype") else ""
    details = (
        "### Human Fix Approval Required\n\n"
        f"**Healing attempt:** {used + 1} of {MAX_HEALING}\n\n"
        f"**Failure analysis:** {analysis.get('failure_type', '?')}{subtype} — {analysis.get('root_cause', 'n/a')}\n\n"
        f"**Risky:** {payload.get('risky')}\n\n"
        f"**Reason:** {payload.get('reason')}\n\n"
        "**Proposed Patch:**\n\n```diff\n" f"{payload.get('patch_diff', '')}\n```"
    )
    return gr.update(value=details, visible=True), gr.update(visible=True)


def _view(thread_id: str | None, status: str, session=NOOP, history=NOOP):
    """One UI frame: [status, session, progress, approval_details, approval_buttons, return_btn, history]."""
    if thread_id is None:
        return status, session, NOOP, NOOP, NOOP, NOOP, history
    details, buttons = approval_ui({"thread_id": thread_id})
    return status, session, _progress_html(thread_id), details, buttons, gr.update(visible=True), history


def start_graph_workflow(url, requirements, browser, mode, maximum, timeout, authorization, use_mock,
                         mock_outcome, mock_risky, mock_fixes=1, llm_cfg=None, history=None,
                         auth_username="", auth_password="", skip_auth=False):
    """Start the LangGraph workflow and stream real node progress to the UI."""
    history = list(history or [])
    problems = []
    if not (url or "").strip():
        problems.append("Enter an application URL.")
    elif not re.match(r"^https?://[^\s/]+", url.strip(), re.I):
        problems.append("Enter a complete URL beginning with https:// or http://.")
    if not (requirements or "").strip():
        problems.append("Describe at least one test requirement.")
    if not authorization:
        problems.append("Confirm that you are authorized to test this application.")
    username, password, skip_auth = (auth_username or "").strip(), auth_password or "", bool(skip_auth)
    if bool(username) != bool(password) and not skip_auth:
        problems.append("Enter both a username and a password, or leave both empty.")
    if problems:
        yield _view(None, "### Please correct these items\n" + "\n".join(f"- {p}" for p in problems))
        return

    llm_public, model_string, settings = None, None, None
    if not use_mock:
        cfg = llm_cfg or {}
        if not cfg.get("validated"):
            yield _view(None, "### Live mode needs validated LLM settings\nOpen **Settings → LLM configuration** and "
                              "save valid credentials, or enable **Use mock mode** for a simulated run. "
                              "No mock data is used in live mode.")
            return
        settings = LLMSettings(cfg["provider"], cfg["model"], cfg.get("api_key", ""), True)
        problem = check_static(settings.provider, settings.model, settings.api_key)
        if problem:
            yield _view(None, f"### LLM settings unusable\n{problem}")
            return
        model_string, llm_public = settings.model_string, settings.public()

    thread_id = str(uuid4())
    auth, auth_skipped = AuthConfig(), False
    if not use_mock:
        # Preflight: deterministic checks BEFORE the workflow, so a wrong URL or login never costs LLM tokens.
        candidate = AuthConfig()
        try:
            if username and password and not skip_auth:
                candidate = build_auth("username_password", {"username": username, "password": password})
                register_run_auth(thread_id, candidate)  # lets redact() scrub the values while Preflight runs
            yield _view(None, "### Preflight\nChecking that the website is reachable and whether it needs a login…")
            result = run_preflight(url.strip(), username if candidate.required else "",
                                   password if candidate.required else "", skip_auth=skip_auth)
        except BaseException:
            release_run_auth(thread_id)
            raise
        if not (result.passed and result.auth_required):
            release_run_auth(thread_id)  # failed, inconclusive, public or skipped: nothing to keep
        if not result.passed:
            yield _view(None, result.markdown())
            return
        auth, auth_skipped = (candidate if result.auth_required else AuthConfig()), result.auth_skipped
    info = {"thread_id": thread_id, "tracker": StageTracker(), "trail": [], "error": None, "paused": False,
            "mock": bool(use_mock), "started": datetime.now(timezone.utc).isoformat(timespec="seconds"),
            "ended": None, "lock": threading.Lock(), "llm": llm_public,
            "secret": effective_key(settings.provider, settings.api_key) if settings else "",
            "meta": {"url": url.strip(), "requirements": requirements.strip(), "browser": browser,
                     "execution_mode": mode, "timeout_seconds": int(timeout)}}
    with _RUNS_LOCK:
        RUNS[thread_id] = info
    history.append(thread_id)
    session = {"thread_id": thread_id, "waiting_for_approval": False}

    request = {"url": url.strip(), "requirements": requirements.strip(), "scenario_cap": int(maximum),
               "authorization": bool(authorization),
               **public_request_fields(auth, auth_skipped)}  # flags only, never the values
    initial_state = {"test_request": request, "healing_retry_count": 0, "errors": []}
    if use_mock:
        initial_state.update(mock_mode=True, mock_execution_outcome=mock_outcome, mock_risky=bool(mock_risky),
                             mock_fixes_needed=int(mock_fixes))
    else:
        request.update(planner_model=model_string, codegen_model=model_string, healer_model=model_string)
        initial_state["mock_mode"] = False
        register_run_credentials(thread_id, settings)
        register_run_auth(thread_id, auth)

    try:
        info["tracker"].begin("intake")
        yield _view(thread_id, "Workflow running…", session, history)
        for _ in _advance(thread_id, initial_state):
            yield _view(thread_id, "Workflow running…", session, history)
        session["waiting_for_approval"] = bool(info["paused"])
        yield _view(thread_id, _status_text(thread_id), session, history)
    except BaseException:
        release_run_secrets(thread_id)  # cancellation or an unexpected error must never leave secrets behind
        raise


def resume_graph_workflow(session, decision):
    """Resume a paused run from its checkpoint with the human decision."""
    thread_id = (session or {}).get("thread_id")
    if thread_id not in RUNS or not _pending_interrupt(thread_id):
        yield _view(None, "No pending approval request.", session)
        return
    if decision not in ("approved", "rejected"):
        yield _view(thread_id, "Invalid approval decision. Choose Approve or Reject.", session)
        return
    info = RUNS[thread_id]
    if not info["lock"].acquire(blocking=False):
        yield _view(thread_id, "A decision for this run is already being processed.", session)
        return
    try:
        info.setdefault("decisions", []).append(decision)
        info["paused"] = False
        yield _view(thread_id, f"Fix {decision}; resuming workflow…", session)
        for _ in _advance(thread_id, Command(resume=decision)):
            yield _view(thread_id, "Workflow running…", session)
        session = {**session, "waiting_for_approval": bool(info["paused"])}
        yield _view(thread_id, _status_text(thread_id), session)
    except BaseException:
        release_run_secrets(thread_id)  # cancelled mid-resume: the run is abandoned, so its secrets go with it
        raise
    finally:
        info["lock"].release()


def approve_workflow(session):
    yield from resume_graph_workflow(session, "approved")


def reject_workflow(session):
    yield from resume_graph_workflow(session, "rejected")


def select_run(thread_id):
    """Switch the active run (each run keeps its own checkpointed state)."""
    if not thread_id or thread_id not in RUNS:
        return NOOP, NOOP, NOOP, NOOP, NOOP
    session = {"thread_id": thread_id, "waiting_for_approval": bool(_pending_interrupt(thread_id))}
    details, buttons = approval_ui(session)
    return session, _status_text(thread_id), _progress_html(thread_id), details, buttons


# ------------------------------------------------------------------- results from real state
def _step_text(step: dict) -> str:
    return " ".join(str(v) for v in (step.get("action"), step.get("target"), step.get("value")) if v)


def build_run(thread_id: str) -> dict | None:
    """Assemble the results view model strictly from the run's checkpointed graph state."""
    info = RUNS.get(thread_id)
    if info is None:
        return None
    values = _snapshot_values(thread_id)
    plan = values.get("test_plan") or []
    result = values.get("execution_result") or {}
    simulated = bool(result.get("simulated")) or info["mock"]
    tag = " (simulated)" if simulated else ""
    tests = {t.get("scenario_name"): t for t in result.get("tests", [])}
    evidence = {e.get("scenario_name"): e for e in values.get("execution_evidence") or []}
    analysis, fix = values.get("failure_analysis"), values.get("fix_proposal")
    healing = values.get("healing_retry_count", 0)
    decision = values.get("approval_decision")
    terminal = values.get("terminal_status") or ("RUNNING" if not info.get("ended") else "UNKNOWN")

    analysis_text = None
    if analysis:
        subtype = f" / {analysis['subtype']}" if analysis.get("subtype") else ""
        analysis_text = f"{analysis['failure_type']}{subtype}: {analysis['root_cause']}\n\nEvidence: {analysis['evidence']}"
    fix_text = None
    if fix:
        fix_text = (f"{fix['reason']}  \nRisky: {fix['risky']} · Decision: {decision or 'pending'} · "
                    f"Applied healing attempts: {healing} of {MAX_HEALING}\n\n```diff\n{fix['patch_diff']}\n```")

    rows, total_ms = [], 0
    for plan_item in plan:
        name = plan_item.get("name", "Unnamed scenario")
        test, ev = tests.get(name), evidence.get(name) or {}
        if test:
            passed = test.get("status") == "passed"
            status, ms = ("Passed" if passed else "Failed"), int(test.get("duration_ms") or 0)
            total_ms += ms
        else:
            passed, status, ms = False, ("Blocked" if terminal in ("BLOCKED", "ERROR") else "Not run"), 0
        err = ev.get("error_message")
        evidence_lines = []
        if test and not passed:
            evidence_lines.append(f"Failed step: {ev.get('failed_step') or 'unknown'}")
        for label, key in (("Screenshot", "screenshot_path"), ("Trace", "trace_path"), ("Final URL", "final_url")):
            evidence_lines.append(f"{label}: {ev.get(key) or 'not captured'}")
        evidence_lines.append(f"Console entries: {len(ev.get('console_log') or [])} · "
                              f"Network errors: {len(ev.get('network_errors') or [])}")
        healed = passed and healing > 0
        if passed:
            actual = f"Passed{tag}" + (f" after {healing} healing attempt(s)" if healed else "")
        elif test:
            actual = f"Failed{tag}: {err or 'no error message'}"
        else:
            actual = "Not executed."
        rows.append({
            "Scenario ID": f"S-{len(rows) + 1:03d}", "Scenario Name": name, "Status": status,
            "Duration": f"{ms / 1000:.1f}s{tag}" if test else "—",
            "Error Summary": (err or "—") if test else ("Not executed" if status == "Not run" else f"Not executed ({terminal})"),
            "Test Steps": [_step_text(s) for s in plan_item.get("steps", [])] or [name],
            "Expected Result": plan_item.get("expected_result") or "Not specified",
            "Actual Result": actual,
            "Error Message": err or ("None" if test else "Not executed."),
            "Evidence Path": "\n".join(evidence_lines) if test else "No evidence collected.",
            "Analysis": (("Previously failed. " if healed else "") + analysis_text) if (analysis_text and (not passed or healed))
                        else ("No failure to analyse." if passed else "No analysis available."),
            "Proposed Fix": fix_text if (fix_text and (not passed or healed)) else ("Not needed." if passed else "Not generated."),
        })
    failed = sum(r["Status"] == "Failed" for r in rows)
    passed_n = sum(r["Status"] == "Passed" for r in rows)
    return {
        "run_id": values.get("run_id") or thread_id, "thread_id": thread_id, "created_at": info["started"],
        "mode": "Mock mode — simulated agents and browser" if info["mock"] else "Live",
        "simulated": info["mock"], "terminal_status": terminal, "healing_attempts": healing,
        "models_used": values.get("models_used") or {}, "errors": list(values.get("errors") or []),
        "trail": list(info["trail"]), "workflow_error": info.get("error"), "paused": bool(info.get("paused")),
        "llm": info.get("llm"), **info["meta"], "scenarios": rows,
        "passed": passed_n, "failed": failed, "skipped": len(rows) - passed_n - failed,
        "execution_time": f"{total_ms / 1000:.1f}s" if tests else "Not executed",
        "stages": dict(info["tracker"].status),
    }


def run_summary_markdown(run: dict | None) -> str:
    if not run:
        return "No run yet. Start a test from **User Request**."
    models = ", ".join(f"{k}: {v}" for k, v in run["models_used"].items()) or "none recorded"
    if run["simulated"]:
        llm = "Mock mode — no LLM provider is called; LLM settings are not used."
    else:
        llm = f"{run['llm']['provider']} · {run['llm']['model']} (key not shown)" if run["llm"] else "n/a"
    lines = [f"**Result:** {run['terminal_status']}" + (" — **awaiting your approval on the User Request page**" if run["paused"] else ""),
             f"**Mode:** {run['mode']}", f"**LLM:** {llm}", f"**Models used by agents:** {models}",
             f"**Healing attempts:** {run['healing_attempts']} of {MAX_HEALING}",
             f"**Executed nodes:** {' → '.join(run['trail']) or 'none'}"]
    if run["workflow_error"]:
        lines.append(f"**Workflow error:** {run['workflow_error']}")
    if run["errors"]:
        lines.append("**Recorded errors:** " + "; ".join(run["errors"]))
    return "  \n".join(lines)


def refresh_results(session, history):
    """Rebuild every Results & Reports component from the active run's graph state."""
    thread_id = (session or {}).get("thread_id")
    run = build_run(thread_id) if thread_id else None
    labels = ["Total Scenarios", "Passed", "Failed", "Skipped",
              "Execution Time (simulated)" if run and run["simulated"] else "Execution Time", "Healing Attempts"]
    values = ([str(len(run["scenarios"])), str(run["passed"]), str(run["failed"]), str(run["skipped"]),
               run["execution_time"], f"{run['healing_attempts']} / {MAX_HEALING}"] if run else ["—"] * 6)
    cards = [gr.update(value=metric_card(v, lbl)) for v, lbl in zip(values, labels)]
    table = [[s["Scenario ID"], s["Scenario Name"], s["Status"], s["Duration"], s["Error Summary"]]
             for s in run["scenarios"]] if run else []
    choices = [(f"{RUNS[t]['started'][11:19]} · {'Mock' if RUNS[t]['mock'] else 'Live'} · {RUNS[t]['meta']['url']}", t)
               for t in (history or []) if t in RUNS]
    return (run, *cards, gr.update(value=table), "Select a scenario row to inspect its details.",
            run_summary_markdown(run), gr.update(choices=choices, value=thread_id))


# ------------------------------------------------------------------- page renderers (read-only views of real state)
_DIFF_RE = re.compile(r"```diff\n(.*?)\n```", re.S)


def _tid(session) -> str | None:
    thread_id = session.get("thread_id") if isinstance(session, dict) else None
    return thread_id if thread_id in RUNS else None


def _pending_count() -> int:
    return sum(1 for info in RUNS.values() if info.get("paused"))


def _text_html(text: object) -> str:
    return html.escape(str(text if text is not None else "")).replace("\n", "<br>")


def workflow_graph_html(thread_id: str | None) -> str:
    info = RUNS.get(thread_id) if thread_id else None
    if info is None:
        head = D.note("No run selected. The graph shows the real LangGraph structure; node colours appear once a run starts.", "info")
        svg = D.workflow_svg(None)
    else:
        tracker = info["tracker"]
        if info.get("error"):
            state = D.badge("Stopped", "bad")
        elif info.get("paused"):
            state = D.badge("Waiting for approval", "warn")
        elif info.get("ended"):
            state = D.badge("Finished", "ok")
        else:
            state = D.badge("Running", "info")
        mode = D.badge("Mock", "violet") if info["mock"] else D.badge("Live", "info")
        head = (f"<div class='tag-row' style='margin:0 0 6px'>{state}{mode}<span class='muted' style='margin:0'>"
                f"{html.escape(info['meta']['url'])}</span></div>")
        status, notes = tracker.view()
        svg = D.workflow_svg(status, notes, tracker.count)
    return f"{head}<div class='wf-card'>{svg}</div>{D.workflow_legend()}"


def workflow_details_html(thread_id: str | None) -> str:
    info = RUNS.get(thread_id) if thread_id else None
    if info is None:
        trail = D.empty("No execution trail yet", "The nodes that actually ran are listed here for the selected run.", "workflow")
    else:
        steps = "".join(f"<div><b>{i}.</b> {html.escape(step)}</div>" for i, step in enumerate(info["trail"], 1))
        trail = f"<div class='log'>{steps or 'Waiting for the first node to finish…'}</div>"
        if info.get("error"):
            trail += D.note(_text_html(info["error"]).replace("**", ""), "bad")
        passes = [h for h in info["tracker"].history if info["tracker"].count[h["stage"]] > 1]
        if passes:
            trail += "<div style='margin-top:12px'>" + D.table(
                ["Stage", "Pass", "Status", "Note"], [[h["stage"], f"#{h['pass']}", h["status"], h["note"] or "—"] for h in passes]) + "</div>"
    return (f"<div class='stack'>{D.panel('Execution trail', trail, 'Nodes that really ran, in order')}"
            f"{D.panel('Routing rules', D.routing_table(), 'Mirrors graph/routing.py: what decides each branch')}</div>")


def approval_card_html(thread_id: str | None) -> str:
    info = RUNS.get(thread_id) if thread_id else None
    payload = _pending_interrupt(thread_id) if info else None
    values = _snapshot_values(thread_id) if info else {}
    analysis = values.get("failure_analysis") or {}
    if not payload:
        fix, decisions = values.get("fix_proposal"), (info or {}).get("decisions") or []
        if info and fix and decisions:
            return (f"<div class='panel'><div class='panel-head'><div class='panel-title'>Last decision for this run</div>"
                    f"{D.badge(decisions[-1])}</div><div class='panel-body stack'><p class='muted'>{html.escape(str(fix.get('reason', '')))}</p>"
                    f"{D.diff_html(fix.get('patch_diff', ''))}</div></div>")
        return D.empty("No fix is awaiting approval",
                       "When a failed test gets a proposed fix, it appears here and the workflow waits. Nothing is ever applied automatically.", "shield")
    evidence = next((e for e in values.get("execution_evidence") or [] if e.get("error_message")), {})
    subtype = f" / {analysis['subtype']}" if analysis.get("subtype") else ""
    confidence = str(analysis.get("confidence", "n/a"))
    tone = {"high": "ok", "medium": "warn", "low": "bad"}.get(confidence, "mute")
    used = values.get("healing_retry_count", 0)
    risky = D.badge("Risky fix", "bad") if payload.get("risky") else D.badge("Not flagged risky", "ok")
    facts = (f"<dl class='kv'><dt>Identified issue</dt><dd>{html.escape(str(evidence.get('scenario_name') or '—'))}: "
             f"{html.escape(str(evidence.get('error_message') or 'no error message recorded'))}</dd>"
             f"<dt>Failure class</dt><dd>{html.escape(str(analysis.get('failure_type', '?')))}{html.escape(subtype)}</dd>"
             f"<dt>Root cause</dt><dd>{html.escape(str(analysis.get('root_cause', 'n/a')))}</dd>"
             f"<dt>Confidence</dt><dd>{D.badge(confidence, tone)}</dd>"
             f"<dt>Suggested fix</dt><dd>{html.escape(str(payload.get('reason', '')))}</dd>"
             f"<dt>Healing attempt</dt><dd>{used + 1} of {MAX_HEALING}</dd>"
             f"<dt>Run mode</dt><dd>{D.badge('Mock', 'violet') if info['mock'] else D.badge('Live', 'info')}</dd></dl>")
    warn = D.note("The Fix Proposal agent flagged this patch as <b>risky</b>. Review it carefully before deciding.", "warn") if payload.get("risky") else ""
    guide = D.note("<b>Approve</b> applies this patch, then the workflow re-validates and re-runs the test. "
                   "<b>Reject</b> discards it and finishes the run with its current result.", "info")
    return (f"<div class='panel approval-box'><div class='panel-head'><div><div class='panel-title'>Human fix approval required</div>"
            f"<div class='panel-sub'>The workflow is paused at the approval gate. It resumes only with your decision.</div></div>"
            f"<div class='tag-row'>{D.badge('Waiting for approval', 'warn')}{risky}</div></div><div class='panel-body stack'>"
            f"{warn}{facts}<div class='panel-title'>Proposed patch</div>{D.diff_html(payload.get('patch_diff', ''))}{guide}</div></div>")


def approval_history_html() -> str:
    rows = []
    for rec in sorted((r for r in H.load_records() if r.get("source") == "ui" and r.get("decisions")),
                      key=lambda r: r.get("started") or "", reverse=True)[:15]:
        rows.append([V._fmt_time(rec.get("started")), V._short_url(rec.get("url") or "—"), V._mode_badge(rec),
                     " ".join(D.badge(d) for d in rec["decisions"]), D.badge(rec.get("status"))])
    if not rows:
        return D.empty("No decisions recorded yet", "Approvals and rejections are logged here with the run's final status.", "shield")
    return D.table(["Run started", "Application", "Mode", "Decisions (in order)", "Final status"], rows, raw_cols=(2, 3, 4))


def results_extras_html(thread_id: str | None) -> tuple[str, str, str]:
    run = build_run(thread_id) if thread_id in RUNS else None
    if not run:
        empty = D.empty("No run selected", "Start a test or choose a run above.", "inbox")
        return empty, empty, empty
    values = _snapshot_values(thread_id)
    cards = []
    analysis = values.get("failure_analysis")
    if analysis:
        cards.append(D.note(f"Analyzer: <b>{html.escape(str(analysis.get('failure_type')))}</b>"
                            f"{html.escape(' / ' + analysis['subtype']) if analysis.get('subtype') else ''} · confidence "
                            f"<b>{html.escape(str(analysis.get('confidence', 'n/a')))}</b>", "info"))
    for sc in run["scenarios"]:
        fix = sc["Proposed Fix"]
        m = _DIFF_RE.search(fix)
        fix_html = (_text_html(_DIFF_RE.sub("", fix).strip()) + D.diff_html(m.group(1))) if m else _text_html(fix)
        body = (f"<dl class='kv'><dt>Error message</dt><dd>{_text_html(sc['Error Message'])}</dd>"
                f"<dt>Analysis</dt><dd>{_text_html(sc['Analysis'])}</dd><dt>Proposed fix</dt><dd>{fix_html}</dd></dl>")
        cards.append(D.panel(f"{sc['Scenario ID']} · {sc['Scenario Name']}", body, right=D.badge(sc["Status"])))
    failure_html = f"<div class='stack'>{''.join(cards)}</div>" if run["scenarios"] else D.empty("No scenarios", "This run produced no test plan.", "inbox")

    rows = [[e.get("scenario_name") or "—", e.get("screenshot_path") or "not captured", e.get("trace_path") or "not captured",
             e.get("final_url") or "—", str(len(e.get("console_log") or [])), str(len(e.get("network_errors") or []))]
            for e in values.get("execution_evidence") or []]
    artifacts = (D.table(["Scenario", "Screenshot", "Trace", "Final URL", "Console", "Net errors"], rows)
                 + (D.note("Mock mode: these artifacts are simulated.", "violet") if run["simulated"] else "")) if rows else \
        D.empty("No artifacts collected", "Evidence (screenshots, traces, console and network logs) appears here after execution.", "inbox")

    steps = "".join(f"<div><b>{i}.</b> {html.escape(step)}</div>" for i, step in enumerate(run["trail"], 1))
    log = f"<div class='log'>{steps or 'No nodes have finished yet.'}</div>"
    if run["errors"]:
        log += D.note("Recorded errors: " + html.escape("; ".join(run["errors"])), "bad")
    if run["workflow_error"]:
        log += D.note(_text_html(run["workflow_error"]).replace("**", ""), "bad")
    return failure_html, artifacts, f"<div class='panel'><div class='panel-body'>{log}</div></div>"


def report_overview_html(thread_id: str | None) -> str:
    run = build_run(thread_id) if thread_id in RUNS else None
    if not run:
        return D.empty("No run selected", "Select or start a run to see its report summary. Saved report files are listed below.", "file")
    facts = "".join(f"<dt>{html.escape(k)}</dt><dd>{html.escape(str(v))}</dd>" for k, v in _report_facts(run))
    scen = D.table(["ID", "Scenario", "Status", "Duration", "Error summary"],
                   [[s["Scenario ID"], s["Scenario Name"], D.badge(s["Status"]), s["Duration"], s["Error Summary"]] for s in run["scenarios"]],
                   raw_cols=(2,)) if run["scenarios"] else D.empty("No scenarios", "This run produced no test plan.", "inbox")
    decisions = RUNS[thread_id].get("decisions") or []
    outcome = (" ".join(D.badge(d) for d in decisions)) if decisions else "<span class='muted'>No approval decision was needed or recorded.</span>"
    summary = D.panel("Execution summary", f"<dl class='kv'>{facts}</dl>")
    results = f"<div class='stack'>{D.panel('Test results', scen)}{D.panel('Approval history for this run', outcome)}</div>"
    return (D.note(html.escape(_report_notice(run)), "violet" if run["simulated"] else "info")
            + f"<div class='stack'>{summary}{results}</div>")


def _report_frame(page: Path | None, empty_text: str) -> str:
    """Show a saved report.html inside the page. sandbox="" blocks scripts and navigation from the report."""
    if page is None:
        return D.empty("No report to show", empty_text, "file")
    document = page.read_text(encoding="utf-8", errors="replace")[:2_000_000]
    return ("<iframe class='report-frame' sandbox='' title='Test run report' "
            f"srcdoc=\"{html.escape(document, quote=True)}\"></iframe>")


def preview_html(run: dict | None) -> str:
    """The styled HTML report (runs/<run id>/report.html) for the selected run."""
    if not run:
        return _report_frame(None, "Start a test or choose a run to see its report here.")
    return _report_frame(H.report_file(str(run.get("run_id") or ""), "report.html"),
                         "This run has no saved report yet. It is written when the workflow reaches the report step.")


def open_persisted_report(run_dir):
    md, page = H.report_file(run_dir, "report.md"), H.report_file(run_dir, "report.html")
    if not (md or page):
        return _report_frame(None, "Choose a report folder to open its report."), None
    return _report_frame(page, "This run folder has no report.html."), [str(p) for p in (md, page) if p]


def _saved_reports_panel() -> str:
    return D.panel("Saved report files", V.persisted_reports_table(), "Existing runs/<run id>/report.md and report.html files")


def refresh_pages(session):
    """Re-render every read-only page for the active run from real state and persisted history."""
    thread_id = _tid(session)
    pending = _pending_count()
    failure_html, artifacts_html, log_html = results_extras_html(thread_id)
    awaiting = bool(thread_id and _pending_interrupt(thread_id))
    return [V.overview_html(H.load_records(), pending), workflow_graph_html(thread_id), workflow_details_html(thread_id),
            approval_card_html(thread_id), gr.update(visible=awaiting), failure_html, artifacts_html, log_html,
            report_overview_html(thread_id), _saved_reports_panel(), gr.update(choices=H.list_report_dirs()),
            approval_history_html(), gr.update(visible=pending > 0)]


def goto_approval(session):
    thread_id = _tid(session)
    if not (thread_id and RUNS[thread_id].get("paused")):
        thread_id = next((t for t, info in RUNS.items() if info.get("paused")), thread_id)
    if thread_id:
        session = {"thread_id": thread_id, "waiting_for_approval": bool(_pending_interrupt(thread_id))}
    return "Fix Approval", session


def _ui_frames(frames):
    """Map the workflow's UI frames onto the redesigned components (rich approval card + live graph)."""
    for f in frames:
        thread_id = _tid(f[1])
        yield (f[0], f[1], f[2], approval_card_html(thread_id) if thread_id else NOOP, f[4], f[5], f[6],
               workflow_graph_html(thread_id) if thread_id else NOOP)


def start_ui(*args):
    yield from _ui_frames(start_graph_workflow(*args))


def confirmed_decision(session, decision):
    """Runs only after the explicit confirmation step; the real gate (pending-interrupt check) still applies."""
    hide = gr.update(visible=False)
    for frame in _ui_frames(resume_graph_workflow(session, decision)):
        yield (*frame, hide)


def ask_confirm(session, decision):
    thread_id = _tid(session)
    if not thread_id or not _pending_interrupt(thread_id):
        return None, gr.update(visible=False), "", NOOP
    if decision == "approved":
        text, tone = "You are about to <b>approve</b> this patch. It will be applied, then the workflow re-validates and re-runs the test.", "warn"
    else:
        text, tone = "You are about to <b>reject</b> this patch. It will not be applied and the run finishes with its current result.", "bad"
    return decision, gr.update(visible=True), D.note(text, tone), gr.update(visible=False)


def cancel_confirm():
    return None, gr.update(visible=False), gr.update(visible=True)


def select_run_ui(thread_id):
    """Switch the active run; returns [session, status, progress, buttons]."""
    session, status, progress, _details, buttons = select_run(thread_id)
    return session, status, progress, buttons


def build_ui() -> gr.Blocks:
    with gr.Blocks(title="Bugle Automation · AI Test Platform") as demo:
        theme_state = gr.State("Dark")
        run_state = gr.State(None)
        run_history = gr.State([])
        sidebar_collapsed = gr.State(False)
        workflow_session = gr.State(None)
        pending_decision = gr.State(None)
        llm_state = gr.State({})
        theme_style = gr.HTML(theme_css("Dark"), elem_classes="style-holder")
        collapse_style = gr.HTML("", elem_classes="style-holder")
        layout_style = gr.HTML(request_layout_css(True), elem_classes="style-holder")

        with gr.Row(elem_id="app-layout"):
            with gr.Column(elem_id="app-sidebar", scale=0):
                gr.HTML(D.brand_html())
                sidebar_toggle = gr.Button("◀ Collapse", elem_id="sidebar-toggle", size="sm")
                nav = gr.Radio(choices=NAV_CHOICES, value="Workflow", show_label=False, container=False,
                               elem_classes="sidebar-nav", elem_id="sidebar-navigation")
                session_mode = gr.HTML(session_mode_label(True))

            with gr.Column(elem_classes="app-shell"):
                with gr.Row(elem_classes="topbar"):
                    topbar = gr.HTML(topbar_html("Workflow", True))
                    goto_btn = gr.Button("Review pending approval", visible=False, variant="primary", size="sm", scale=0)
                    clear_btn = gr.Button("Clear history", size="sm", scale=0, elem_id="clear-history-btn")
                with gr.Row(visible=False, elem_classes="card", elem_id="clear-confirm") as clear_confirm:
                    gr.HTML("<b>Delete all history?</b> This removes the run log, saved report files and exported reports. "
                            "Test specs and other run artifacts are kept. It cannot be undone.")
                    clear_yes = gr.Button("Yes, delete history", variant="stop", size="sm", scale=0, elem_id="clear-yes")
                    clear_no = gr.Button("Cancel", size="sm", scale=0, elem_id="clear-no")

                # ------------------------------------------------------------------ 1 Overview
                with gr.Column(visible=False, elem_classes="page-col") as overview_page:
                    gr.HTML(D.page_head("Dashboard", "Everything the platform has recorded, at a glance. Numbers come from real runs only.", "overview"))
                    with gr.Row():
                        ov_new = gr.Button("Start a new test", variant="primary", scale=0, min_width=170)
                        ov_analytics = gr.Button("View analytics", scale=0, min_width=150)
                    overview_out = gr.HTML(V.overview_html([], 0))

                # ------------------------------------------------------------- 2 Create & Run
                with gr.Column(visible=False, elem_classes="page-col", elem_id="user-request-page") as request_page:
                    gr.HTML(D.page_head("Create & Run", "Describe the application and behaviours to test. The LangGraph workflow plans, generates, validates and executes the tests.", "play"))
                    with gr.Row(elem_id="request-layout"):
                        with gr.Column(elem_id="request-card-column"):
                            with gr.Group(elem_classes="card"):
                                gr.HTML("<div class='card-title'>Test configuration</div><p class='muted'>Tell us which application and behaviours to cover.</p>")
                                url = gr.Textbox(label="Application URL", placeholder="https://your-app.example",
                                                 info="Full address, starting with https:// or http://.")
                                requirements = gr.Textbox(label="Test requirements", lines=3,
                                                          placeholder="Describe the user journey or behaviours to verify…",
                                                          info="Plain language. At least one requirement is needed.")
                                with gr.Row(elem_id="scenario-controls", equal_height=False):
                                    browser = gr.Dropdown(["Chromium", "Firefox", "WebKit"], value="Chromium", label="Browser", elem_id="browser-control")
                                    maximum = gr.Slider(1, 5, value=3, step=1, label="Maximum number of scenarios")
                                    mode = gr.Radio(["Headless", "Visible Browser"], value="Headless", label="Execution mode", elem_id="execution-mode-control")
                                    timeout = gr.Slider(10, 180, value=30, step=5, label="Timeout (seconds)")
                                with gr.Row(equal_height=False):
                                    auth_username = gr.Textbox(label="Username (optional)", info="Only needed if the website asks for a login.")
                                    auth_password = gr.Textbox(label="Password (optional)", type="password", info="Kept in memory for this run only.")
                                skip_auth = gr.Checkbox(label="I'm testing the login page itself (skip authentication validation)",
                                                        info="The URL must still be reachable. No login is attempted and the site is not treated as signed in.")
                                authorization = gr.Checkbox(label="I am authorized to test this application.")
                                start = gr.Button("Start Testing", variant="primary", size="lg", elem_id="start-testing")
                                request_status = gr.Markdown()
                                return_btn = gr.Button("Open report", visible=False, elem_id="return-to-results")
                        with gr.Column(visible=True, elem_id="progress-card-column") as progress_panel:
                            with gr.Group(elem_classes="card"):
                                gr.HTML("<div class='card-title'>Execution progress</div><p class='muted'>Stages update live from the workflow nodes that actually execute, including approval, healing and reruns.</p>")
                                progress = gr.HTML(render_progress())
                                gr.HTML("<p class='muted'>Mock-mode runs use simulated agents and a simulated browser. Live runs use the model selected in Settings and stop with a clear message at any stage that is not implemented yet.</p>")

                # --------------------------------------------------------------- 3 AI Workflow
                with gr.Column(visible=True, elem_classes="page-col") as workflow_page:
                    gr.HTML(D.page_head("Workflow", "How Bugle tests an application, step by step: plan, write, check, run, explain the failure, fix, report. Colours update live during a run.", "workflow"))
                    with gr.Group(elem_classes="card"):
                        wf_graph = gr.HTML(workflow_graph_html(None))
                    wf_details = gr.HTML(workflow_details_html(None))

                # ------------------------------------------------------------ 5 Fix & Approval
                with gr.Column(visible=False, elem_classes="page-col") as approval_page:
                    gr.HTML(D.page_head("Fix Approval", "Review the proposed fix. A risky fix is never applied until a human approves it.", "shield"))
                    approval_card = gr.HTML(approval_card_html(None))
                    with gr.Row(visible=False) as approval_buttons:
                        approve_btn = gr.Button("Approve Fix", variant="primary", elem_id="approve-btn")
                        reject_btn = gr.Button("Reject Fix", variant="stop", elem_id="reject-btn")
                    with gr.Column(visible=False, elem_classes="card") as confirm_panel:
                        confirm_html = gr.HTML()
                        with gr.Row():
                            confirm_btn = gr.Button("Confirm decision", variant="primary", elem_id="confirm-btn")
                            cancel_btn = gr.Button("Cancel")
                    gr.HTML(D.page_head("Approval history", "Every decision submitted through the gate.", "clock"))
                    approval_history = gr.HTML(approval_history_html())

                # ------------------------------------------------------------------- Reports (results + exports)
                with gr.Column(visible=False, elem_classes="page-col") as reports_page:
                    gr.HTML(D.page_head("Reports", "Results, failure analysis and exports for every test run. Pick a run, read it, export it.", "file"))
                    with gr.Group(elem_classes="card"):
                        with gr.Row(equal_height=False):
                            run_picker = gr.Dropdown([], label="Run", info="Each run keeps its own results.", interactive=True, scale=4)
                            export_format = gr.Dropdown(["PDF", "Word (.docx)", "Excel (.xlsx)"], label="Export as", value="PDF", scale=1)
                            export_btn = gr.Button("Export", variant="primary", scale=0, min_width=110)
                        run_summary = gr.Markdown("No run yet. Start a test from **Create & Run**.")
                        export_status = gr.Markdown()
                        export_file = gr.File(label="Download report", visible=False)
                    with gr.Row(elem_classes="kpi-row"):
                        metric_total = gr.HTML(metric_card("—", "Total Scenarios"))
                        metric_pass = gr.HTML(metric_card("—", "Passed"))
                        metric_fail = gr.HTML(metric_card("—", "Failed"))
                        metric_skip = gr.HTML(metric_card("—", "Skipped"))
                        metric_time = gr.HTML(metric_card("—", "Execution Time"))
                        metric_heal = gr.HTML(metric_card("—", "Healing Attempts"))
                    with gr.Tabs():
                        with gr.Tab("Report"):
                            with gr.Group(elem_classes="card"):
                                preview_btn = gr.Button("Refresh report")
                                report_preview = gr.HTML(_report_frame(None, "Start a test or choose a run to see its report here."))
                        with gr.Tab("Summary"):
                            report_overview = gr.HTML(report_overview_html(None))
                        with gr.Tab("Scenarios"):
                            with gr.Group(elem_classes="card"):
                                scenarios_table = gr.Dataframe(headers=["Scenario ID", "Scenario Name", "Status", "Duration", "Error Summary"],
                                                               datatype=["str"] * 5, interactive=False, value=[])
                                scenario_info = gr.Markdown("Select a scenario row to inspect its details.")
                        with gr.Tab("Failure analysis"):
                            res_analysis = gr.HTML(results_extras_html(None)[0])
                        with gr.Tab("Evidence"):
                            res_artifacts = gr.HTML(results_extras_html(None)[1])
                        with gr.Tab("Execution log"):
                            res_log = gr.HTML(results_extras_html(None)[2])
                        with gr.Tab("Saved reports"):
                            persisted_table = gr.HTML(_saved_reports_panel())
                            with gr.Group(elem_classes="card"):
                                persisted_pick = gr.Dropdown(H.list_report_dirs(), label="Open a saved report", interactive=True)
                                persisted_open = gr.Button("Open")
                                persisted_md = gr.HTML()
                                persisted_files = gr.File(label="Report files", file_count="multiple")

                # ----------------------------------------------------------------- 7 Analytics
                with gr.Column(visible=False, elem_classes="page-col") as analytics_page:
                    gr.HTML(D.page_head("Analytics & Monitoring", "Trends and distributions calculated from recorded runs. Empty ranges stay empty; nothing is estimated.", "chart"))
                    with gr.Group(elem_classes="card"):
                        with gr.Row():
                            an_window = gr.Radio(["Last 24 hours", "Last 7 days", "Last 30 days", "All time"], value="All time", label="Time range")
                            an_source = gr.Dropdown(["All sources", "UI runs", "Report files"], value="All sources", label="Data source")
                            an_mock = gr.Checkbox(value=True, label="Include mock runs")
                            an_refresh = gr.Button("Refresh", scale=0, min_width=110)
                    analytics_out = gr.HTML(V.analytics_html("All time", "All sources", True))

                # ------------------------------------------------------------------ 8 Settings
                with gr.Column(visible=False, elem_classes="page-col") as settings_page:
                    gr.HTML(D.page_head("Settings", "Configure appearance, mock-mode behaviour, and your LLM connection.", "settings"))
                    with gr.Group(elem_classes="card"):
                        gr.HTML("<div class='card-title'>Appearance</div><p class='muted'>Choose a theme for this active UI session.</p>")
                        settings_theme = gr.Radio(["Dark", "Violet", "Light"], value="Dark", label="Dashboard theme")
                    with gr.Group(elem_classes="card"):
                        gr.HTML("<div class='card-title'>Execution mode</div><p class='muted'>Mock mode runs the real LangGraph workflow with simulated agents (no real LLM or external calls), browser and fixes (clearly labeled). Turn it off for live mode, which uses your LLM settings and never falls back to mocks.</p>")
                        use_mock = gr.Checkbox(value=True, label="Use mock mode (no real LLM)")
                        mock_outcome = gr.Radio(choices=["PASS", "FAIL"], value="PASS", label="Mock execution outcome")
                        mock_risky = gr.Checkbox(value=False, label="Simulate risky fix")
                        mock_fixes = gr.Slider(1, 3, value=1, step=1, label="Mock mode: fixes needed before the rerun passes",
                                               info="Only applies when the outcome is FAIL. 1-2 heal successfully; 3 exhausts the 2-attempt limit and ends FAIL.")
                        show_progress = gr.Checkbox(value=True, label="Show execution progress")
                    with gr.Group(elem_classes="card"):
                        gr.HTML("<div class='card-title'>LLM configuration</div><p class='muted'>Used by live mode only. Credentials are validated with a real provider call, kept for this UI session only, and never shown in reports, exports or messages. Mock-mode runs do not call any provider.</p>")
                        llm_provider = gr.Dropdown(["Groq", "Gemini", "OpenRouter", "OpenAI", "Anthropic", "Ollama"], value="OpenRouter", label="Provider")
                        llm_model = gr.Textbox(label="Model name", placeholder="Enter the provider's model identifier (e.g. gemini-2.0-flash)")
                        llm_api_key = gr.Textbox(label="API key", type="password", placeholder="Masked and session-only; usually optional for Ollama")
                        save_llm_btn = gr.Button("Validate & save LLM settings")
                        llm_status = gr.Markdown()

        page_columns = [workflow_page, overview_page, request_page, approval_page, reports_page, analytics_page, settings_page]
        assert len(page_columns) == len(PAGES)
        page_outputs = [overview_out, wf_graph, wf_details, approval_card, approval_buttons, res_analysis, res_artifacts, res_log,
                        report_overview, persisted_table, persisted_pick, approval_history, goto_btn]
        results_outputs = [run_state, metric_total, metric_pass, metric_fail, metric_skip, metric_time, metric_heal,
                           scenarios_table, scenario_info, run_summary, run_picker]
        run_outputs = [request_status, workflow_session, progress, approval_card, approval_buttons, return_btn, run_history, wf_graph]
        analytics_inputs = [an_window, an_source, an_mock]

        def ask_clear():
            return gr.update(visible=True)

        def do_clear(history):
            with _RUNS_LOCK:
                for tid in [t for t, info in RUNS.items() if info.get("ended") or info.get("paused") or info.get("error")]:
                    RUNS.pop(tid, None)
                    release_run_secrets(tid)
            H.clear_history((EXPORT_DIR,))
            return None, [], gr.update(visible=False), "History cleared."

        def nav_after_decision(session):
            """Once the workflow has resumed and finished, show the report; if another fix awaits approval, stay."""
            thread_id = _tid(session)
            if thread_id and not _pending_interrupt(thread_id):
                return "Reports"
            return NOOP

        def nav_to_approval(session):
            """When a run pauses for human approval, open the Fix Approval page so the decision is not missed."""
            thread_id = _tid(session)
            if thread_id and _pending_interrupt(thread_id):
                return "Fix Approval"
            return NOOP

        def on_nav(page, mock):
            return [*navigate(page), topbar_html(page, mock)]

        def after(event):
            return event.then(refresh_results, inputs=[workflow_session, run_history], outputs=results_outputs
                              ).then(refresh_pages, inputs=workflow_session, outputs=page_outputs
                                     ).then(preview_html, inputs=run_state, outputs=report_preview)

        sidebar_toggle.click(toggle_sidebar, inputs=[sidebar_collapsed, nav], outputs=[sidebar_collapsed, collapse_style, nav, sidebar_toggle])
        nav.change(on_nav, inputs=[nav, use_mock], outputs=[*page_columns, topbar]
                   ).then(refresh_pages, inputs=workflow_session, outputs=page_outputs
                          ).then(V.analytics_html, inputs=analytics_inputs, outputs=analytics_out)
        demo.load(refresh_pages, inputs=workflow_session, outputs=page_outputs)
        ov_new.click(lambda: "Create & Run", outputs=nav)
        ov_analytics.click(lambda: "Analytics", outputs=nav)
        settings_theme.change(set_theme, inputs=settings_theme, outputs=[theme_style, theme_state])
        use_mock.change(session_mode_label, inputs=use_mock, outputs=session_mode)
        use_mock.change(lambda mock, page: topbar_html(page, mock), inputs=[use_mock, nav], outputs=topbar)
        show_progress.change(progress_visibility, inputs=show_progress, outputs=[progress_panel, layout_style])
        save_llm_btn.click(save_llm_settings, inputs=[llm_provider, llm_model, llm_api_key], outputs=[llm_status, llm_state])

        started = start.click(start_ui, inputs=[url, requirements, browser, mode, maximum, timeout, authorization, use_mock,
                                                mock_outcome, mock_risky, mock_fixes, llm_state, run_history,
                                                auth_username, auth_password, skip_auth], outputs=run_outputs)
        after(started.then(nav_to_approval, inputs=workflow_session, outputs=nav))

        # Two-step decision: the buttons only ask for confirmation; only "Confirm" calls the real resume path.
        approve_btn.click(lambda s: ask_confirm(s, "approved"), inputs=workflow_session,
                          outputs=[pending_decision, confirm_panel, confirm_html, approval_buttons])
        reject_btn.click(lambda s: ask_confirm(s, "rejected"), inputs=workflow_session,
                         outputs=[pending_decision, confirm_panel, confirm_html, approval_buttons])
        cancel_btn.click(cancel_confirm, outputs=[pending_decision, confirm_panel, approval_buttons])
        decided = confirm_btn.click(confirmed_decision, inputs=[workflow_session, pending_decision], outputs=[*run_outputs, confirm_panel])
        after(decided.then(nav_after_decision, inputs=workflow_session, outputs=nav))

        clear_btn.click(ask_clear, outputs=clear_confirm)
        clear_no.click(lambda: gr.update(visible=False), outputs=clear_confirm)
        after(clear_yes.click(do_clear, inputs=run_history, outputs=[workflow_session, run_history, clear_confirm, request_status]))

        after(run_picker.input(select_run_ui, inputs=run_picker, outputs=[workflow_session, request_status, progress, approval_buttons]))
        return_btn.click(lambda: "Reports", outputs=nav)
        goto_btn.click(goto_approval, inputs=workflow_session, outputs=[nav, workflow_session]
                       ).then(refresh_pages, inputs=workflow_session, outputs=page_outputs)
        scenarios_table.select(scenario_details, inputs=run_state, outputs=scenario_info)
        preview_btn.click(preview_html, inputs=run_state, outputs=report_preview)
        def export_ui(run, fmt):
            status, path = export_report(run, fmt)
            return status, gr.update(value=path, visible=bool(path))

        export_btn.click(export_ui, inputs=[run_state, export_format], outputs=[export_status, export_file])
        persisted_open.click(open_persisted_report, inputs=persisted_pick, outputs=[persisted_md, persisted_files])
        for control in analytics_inputs:
            control.change(V.analytics_html, inputs=analytics_inputs, outputs=analytics_out)
        an_refresh.click(V.analytics_html, inputs=analytics_inputs, outputs=analytics_out)

    return demo


def launch() -> None:
    build_ui().launch(server_name="127.0.0.1", server_port=int(os.getenv("GRADIO_SERVER_PORT", "7860")), theme=gr.themes.Base(), css=CSS)


if __name__ == "__main__":
    launch()
